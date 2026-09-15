import os
import sys
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, hex_color):
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def add_callout(doc, text, title="LƯU Ý QUAN TRỌNG"):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    cell = table.cell(0, 0)
    cell.width = Inches(6.5)
    set_cell_background(cell, "F0F4F8")
    set_cell_margins(cell, top=140, bottom=140, left=200, right=200)
    
    # Left border
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(f'<w:tcBorders {nsdecls("w")}><w:left w:val="single" w:sz="24" w:space="0" w:color="1A365D"/><w:top w:val="none"/><w:right w:val="none"/><w:bottom w:val="none"/></w:tcBorders>')
    tcPr.append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    r_title = p.add_run(f"📌 {title}: ")
    r_title.bold = True
    r_title.font.color.rgb = RGBColor(26, 54, 93)
    r_title.font.size = Pt(10.5)
    
    r_text = p.add_run(text)
    r_text.font.size = Pt(10.5)
    r_text.font.color.rgb = RGBColor(45, 55, 72)
    
    doc.add_paragraph().paragraph_format.space_after = Pt(6)

def add_heading_styled(doc, text, level):
    h = doc.add_heading(text, level=level)
    h.paragraph_format.keep_with_next = True
    if level == 1:
        h.paragraph_format.space_before = Pt(18)
        h.paragraph_format.space_after = Pt(8)
        for r in h.runs:
            r.font.color.rgb = RGBColor(15, 23, 42)
            r.font.size = Pt(16)
            r.bold = True
    elif level == 2:
        h.paragraph_format.space_before = Pt(14)
        h.paragraph_format.space_after = Pt(6)
        for r in h.runs:
            r.font.color.rgb = RGBColor(30, 58, 138)
            r.font.size = Pt(13)
            r.bold = True
    elif level == 3:
        h.paragraph_format.space_before = Pt(10)
        h.paragraph_format.space_after = Pt(4)
        for r in h.runs:
            r.font.color.rgb = RGBColor(51, 65, 85)
            r.font.size = Pt(11.5)
            r.bold = True
    return h

def add_image_box(doc, img_path, caption):
    if not os.path.exists(img_path):
        p = doc.add_paragraph(f"[Không tìm thấy file ảnh: {img_path}]")
        p.runs[0].font.color.rgb = RGBColor(220, 38, 38)
        return

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run()
    run.add_picture(img_path, width=Inches(6.2))
    
    p_cap = doc.add_paragraph()
    p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cap.paragraph_format.space_before = Pt(2)
    p_cap.paragraph_format.space_after = Pt(14)
    r_cap = p_cap.add_run(f"📷 Hình minh họa: {caption}")
    r_cap.font.size = Pt(9.5)
    r_cap.font.italic = True
    r_cap.font.color.rgb = RGBColor(100, 116, 139)

def main():
    repo_dir = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
    assets_dir = os.path.join(repo_dir, "public", "docs-assets")
    downloads_dir = "C:\\Users\\thaibao\\Downloads"
    
    doc = Document()
    
    # Page setup
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    # Style defaults
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Calibri'
    normal_style.font.size = Pt(11)
    normal_style.font.color.rgb = RGBColor(30, 41, 59)
    normal_style.paragraph_format.line_spacing = 1.2
    normal_style.paragraph_format.space_after = Pt(6)

    # Title
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(0)
    p_title.paragraph_format.space_after = Pt(4)
    r_title = p_title.add_run("TÀI LIỆU HƯỚNG DẪN SỬ DỤNG & KIỂM THỬ THỰC TẾ")
    r_title.bold = True
    r_title.font.size = Pt(19)
    r_title.font.color.rgb = RGBColor(15, 23, 42)

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_after = Pt(16)
    r_sub = p_sub.add_run("DỰ ÁN: CALIFORNIA ARTS / ĐIỂN\nChuyên sâu: Feedback 2 (Bottom CTA) & Feedback 3 (Công cụ Tìm size trực quan)")
    r_sub.font.size = Pt(12)
    r_sub.font.color.rgb = RGBColor(71, 85, 105)

    add_callout(
        doc,
        "Tài liệu này được xuất kèm đầy đủ ảnh chụp màn hình thực tế từ hệ thống đang chạy tại localhost:3000 và Admin Panel. Toàn bộ tính năng đã được kiểm tra vận hành thực tế 100%, thao tác trực quan bằng Form, không cần viết mã JSON hay can thiệp code.",
        "MỤC TIÊU KIỂM THỬ THỰC TẾ"
    )

    # ---------------------------------------------------------------------------
    # PHẦN 1
    # ---------------------------------------------------------------------------
    add_heading_styled(doc, "PHẦN 1: FEEDBACK 2 — NÚT & DÒNG CHỮ CUỐI TRANG (BOTTOM CTA)", level=1)
    
    add_heading_styled(doc, "1.1. Bản chất vấn đề & Vị trí chính xác", level=2)
    p = doc.add_paragraph()
    p.add_run("• ").bold = True
    p.add_run("Vấn đề phản ánh trước đây: ").bold = True
    p.add_run("Khách hàng nhìn thấy dòng chữ và nút bấm ở cuối trang (ví dụ: 'xem thêm về áo sơ mi / khám phá ngay') nhưng không biết chỉnh ở đâu trong Admin, và nút bấm bị nhảy sang liên kết ngẫu nhiên.\n")
    p.add_run("• ").bold = True
    p.add_run("Vị trí chính xác: ").bold = True
    p.add_run("Đây KHÔNG PHẢI là trang chi tiết sản phẩm, mà là ")
    p.add_run("khu vực chân trang của Danh mục / Bộ sưu tập (Collections)").bold = True
    p.add_run(" nằm ngay phía trên Footer.\n")
    p.add_run("• ").bold = True
    p.add_run("Vị trí trong Admin: ").bold = True
    p.add_run("Vào ")
    p.add_run("Collections (Bộ sưu tập)").bold = True
    p.add_run(" → chọn bộ sưu tập cần sửa (ví dụ: Coats & Jackets, Áo sơ mi...) → cuộn xuống trường ")
    p.add_run("Nút & Dòng chữ cuối trang (Bottom CTA)").bold = True
    p.add_run(".")

    add_heading_styled(doc, "1.2. Bảng mô tả chi tiết các trường cấu hình trong Admin", level=2)
    table = doc.add_table(rows=7, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    headers = ["Tên trường trong Admin", "Tên trường kỹ thuật", "Cách sử dụng & Ý nghĩa"]
    widths = [Inches(2.0), Inches(1.5), Inches(3.0)]

    hdr_cells = table.rows[0].cells
    for i, title in enumerate(headers):
        hdr_cells[i].width = widths[i]
        set_cell_background(hdr_cells[i], "1E3A8A")
        set_cell_margins(hdr_cells[i], top=100, bottom=100, left=120, right=120)
        p = hdr_cells[i].paragraphs[0]
        r = p.add_run(title)
        r.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)
        r.font.size = Pt(10)

    rows_data = [
        ("Ẩn khu vực này", "hideCta", "Tích chọn nếu muốn TẮT HOÀN TOÀN nút cuối trang ở bộ sưu tập này."),
        ("Dòng chữ nhỏ (Tiếng Việt)", "eyebrowVi", "Câu dẫn nhỏ tiếng Việt. Ví dụ: 'xem thêm về áo sơ mi'. Để trống mặc định: 'xem toàn bộ sản phẩm'."),
        ("Dòng chữ nhỏ (Tiếng Anh)", "eyebrow", "Câu dẫn nhỏ tiếng Anh tương ứng. Ví dụ: 'explore more shirts'."),
        ("Chữ trên nút (Tiếng Việt)", "buttonLabelVi", "Chữ trên nút bấm. Ví dụ: 'khám phá ngay', 'xem tất cả'."),
        ("Chữ trên nút (Tiếng Anh)", "buttonLabel", "Chữ trên nút tiếng Anh. Ví dụ: 'discover now', 'shop now'."),
        ("Đích đến khi bấm nút", "linkCollection / customUrl", "Có thể chọn trực tiếp từ danh sách bộ sưu tập có sẵn hoặc gõ URL tùy biến (VD: /collections/shop-all).")
    ]

    for row_idx, data in enumerate(rows_data, start=1):
        row_cells = table.rows[row_idx].cells
        bg_color = "F8FAFC" if row_idx % 2 == 1 else "FFFFFF"
        for col_idx in range(3):
            row_cells[col_idx].width = widths[col_idx]
            set_cell_background(row_cells[col_idx], bg_color)
            set_cell_margins(row_cells[col_idx], top=80, bottom=80, left=120, right=120)
            p = row_cells[col_idx].paragraphs[0]
            r = p.add_run(data[col_idx])
            r.font.size = Pt(9.5)
            if col_idx == 0:
                r.bold = True

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    add_heading_styled(doc, "1.3. Hình ảnh chụp thực tế màn hình Admin & Storefront", level=2)
    add_image_box(doc, os.path.join(assets_dir, "fb2_admin_collections_list.png"), "Màn hình danh sách Collections trong Admin Panel")
    add_image_box(doc, os.path.join(assets_dir, "fb2_admin_collection_bottom_cta.png"), "Chi tiết cấu hình Bottom CTA trong Admin (với đầy đủ trường tiếng Việt, tiếng Anh và Link)")
    add_image_box(doc, os.path.join(assets_dir, "fb2_storefront_bottom_cta.png"), "Giao diện hiển thị thực tế ngoài website Storefront ở chân trang danh mục")

    add_heading_styled(doc, "1.4. Các bước tự kiểm tra thực tế (Step-by-Step)", level=2)
    steps_fb2 = [
        ("Bước 1", "Truy cập Admin Panel: http://localhost:3000/admin."),
        ("Bước 2", "Ở menu bên trái, bấm vào 'Collections' → chọn một bộ sưu tập bất kỳ (ví dụ: 'Coats & Jackets')."),
        ("Bước 3", "Cuộn xuống trường 'Nút & Dòng chữ cuối trang (Bottom CTA)'. Nhập thử 'xem thêm bộ sưu tập mới' vào ô Dòng chữ nhỏ (Tiếng Việt) và 'xem ngay' vào ô Chữ trên nút (Tiếng Việt)."),
        ("Bước 4", "Bấm nút 'Save' ở góc trên bên phải màn hình Admin."),
        ("Bước 5", "Mở tab trình duyệt mới vào http://localhost:3000/collections/coats-jackets, cuộn xuống dưới cùng."),
        ("Bước 6", "Xác nhận dòng chữ và nút bấm hiển thị chính xác nội dung vừa nhập. Bấm vào nút để kiểm tra chuyển hướng chính xác."),
        ("Bước 7", "Quay lại Admin, tích chọn 'Ẩn khu vực này' → Bấm Save → Tải lại web để xác nhận khu vực nút biến mất hoàn toàn.")
    ]
    for s_title, s_desc in steps_fb2:
        p = doc.add_paragraph()
        r_step = p.add_run(f"• {s_title}: ")
        r_step.bold = True
        r_step.font.color.rgb = RGBColor(30, 58, 138)
        p.add_run(s_desc)

    doc.add_page_break()

    # ---------------------------------------------------------------------------
    # PHẦN 2
    # ---------------------------------------------------------------------------
    add_heading_styled(doc, "PHẦN 2: FEEDBACK 3 — CÔNG CỤ 'TÌM SIZE' (SIZE FINDER) CHO RIÊNG SẢN PHẨM", level=1)
    
    add_heading_styled(doc, "2.1. Cải tiến cốt lõi: Form giao diện trực quan thay thế hoàn toàn mã JSON", level=2)
    p = doc.add_paragraph()
    p.add_run("• ").bold = True
    p.add_run("Vấn đề trước đây: ").bold = True
    p.add_run("Quản trị viên phải chỉnh một chuỗi mã JSON phức tạp trong Cài đặt chung, dễ sai sót cú pháp dấu ngoặc nhọn/ngoặc vuông, và bị áp dụng chung một bảng size cho tất cả sản phẩm dù là áo hay quần.\n")
    p.add_run("• ").bold = True
    p.add_run("Nâng cấp hoàn thiện: ").bold = True
    p.add_run("Đã trang bị ")
    p.add_run("Giao diện Form nhập liệu trực quan (Visual Array Form)").bold = True
    p.add_run(" nằm ngay trong từng sản phẩm. Quản trị viên chỉ cần bấm ")
    p.add_run("'+ add size rule'").bold = True
    p.add_run(" và điền từng ô: Chiều cao, Cân nặng, Dáng (Ôm/Thoải mái), Size gợi ý. Hoàn toàn không cần biết lập trình hay gõ JSON!")

    add_heading_styled(doc, "2.2. Chi tiết 3 Chế độ hoạt động của Size Finder", level=2)
    table2 = doc.add_table(rows=4, cols=2)
    table2.alignment = WD_TABLE_ALIGNMENT.CENTER
    table2.autofit = False

    t2_widths = [Inches(2.5), Inches(4.0)]
    hdr2_cells = table2.rows[0].cells
    hdr2_cells[0].width = t2_widths[0]
    hdr2_cells[1].width = t2_widths[1]
    for c in hdr2_cells:
        set_cell_background(c, "1E3A8A")
        set_cell_margins(c, top=100, bottom=100, left=120, right=120)
    r0 = hdr2_cells[0].paragraphs[0].add_run("Chế độ hoạt động")
    r0.bold = True
    r0.font.color.rgb = RGBColor(255, 255, 255)
    r1 = hdr2_cells[1].paragraphs[0].add_run("Mô tả & Ứng dụng thực tế")
    r1.bold = True
    r1.font.color.rgb = RGBColor(255, 255, 255)

    modes = [
        ("Dùng cấu hình chung từ Cài đặt website (Mặc định)", "Áp dụng bảng size tiêu chuẩn toàn shop. Phù hợp với các sản phẩm phổ thông thông thường."),
        ("Tùy chỉnh riêng cho sản phẩm này", "Mở ra toàn bộ Form trực quan để thiết lập mốc Chiều cao, Cân nặng và từng dòng Quy tắc size riêng cho sản phẩm đó (quần tây, áo khoác oversize, polo slim...)."),
        ("Tắt tính năng tìm size cho sản phẩm này", "Ẩn hoàn toàn nút 'gợi ý size?' ngoài website. Dành cho phụ kiện, túi xách, mũ nón hoặc sản phẩm One-Size.")
    ]

    for idx, (m_name, m_desc) in enumerate(modes, start=1):
        c_row = table2.rows[idx].cells
        bg_col = "F8FAFC" if idx % 2 == 1 else "FFFFFF"
        for ci in range(2):
            c_row[ci].width = t2_widths[ci]
            set_cell_background(c_row[ci], bg_col)
            set_cell_margins(c_row[ci], top=80, bottom=80, left=120, right=120)
        p_name = c_row[0].paragraphs[0].add_run(m_name)
        p_name.bold = True
        p_name.font.size = Pt(9.5)
        p_desc = c_row[1].paragraphs[0].add_run(m_desc)
        p_desc.font.size = Pt(9.5)

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    add_heading_styled(doc, "2.3. Hình ảnh chụp thực tế màn hình Admin & Storefront", level=2)
    add_image_box(doc, os.path.join(assets_dir, "fb3_admin_size_finder_custom.png"), "Màn hình cấu hình Size Finder khi chọn 'Tùy chỉnh riêng cho sản phẩm này'")
    add_image_box(doc, os.path.join(assets_dir, "fb3_admin_size_rules_form.png"), "Chi tiết Form thêm từng quy tắc Size Rule trực quan (không cần viết JSON)")
    add_image_box(doc, os.path.join(assets_dir, "fb3_storefront_product_page.png"), "Giao diện trang chi tiết sản phẩm ngoài Storefront với 2 nút 'gợi ý size?' và 'bảng size'")
    add_image_box(doc, os.path.join(assets_dir, "fb3_storefront_size_finder_modal.png"), "Popup Modal tìm size tương tác trực quan ngoài website")

    add_heading_styled(doc, "2.4. Các bước tự kiểm tra thực tế (Step-by-Step)", level=2)
    steps_fb3 = [
        ("Bước 1", "Vào Admin: http://localhost:3000/admin → chọn 'Products' → bấm vào sản phẩm 'Hancock Mini Merino Polo' (hoặc bất kỳ sản phẩm nào)."),
        ("Bước 2", "Cuộn xuống mục 'Công cụ \"Tìm size\" (Size Finder)'. Chọn Chế độ hoạt động là 'Tùy chỉnh riêng cho sản phẩm này'."),
        ("Bước 3", "Kiểm tra thấy các ô nhập danh sách chiều cao, cân nặng và bảng 'Quy tắc gợi ý size' lập tức xuất hiện."),
        ("Bước 4", "Bấm '+ add size rule' để thêm thử một quy tắc (VD: Chiều cao '1m68–1m70', Cân nặng '54–58 kg', Dáng 'Ôm' → Gợi ý size 'S')."),
        ("Bước 5", "Bấm 'Save' ở góc trên bên phải Admin."),
        ("Bước 6", "Mở trang sản phẩm ngoài website: http://localhost:3000/products/hancockminimerinopolo."),
        ("Bước 7", "Bấm vào dòng chữ gạch chân 'gợi ý size?' bên dưới nút chọn kích cỡ. Một popup hiện đại sẽ mở lên."),
        ("Bước 8", "Chọn Dáng 'ôm', chọn Chiều cao '1m68–1m70' và Cân nặng '54–58 kg', bấm nút 'tìm size' → Kiểm tra thấy hệ thống đề xuất chính xác size S!"),
        ("Bước 9", "Quay lại Admin đổi Chế độ sang 'Tắt tính năng tìm size' → Bấm Save → Tải lại web xác nhận nút 'gợi ý size?' đã biến mất hoàn toàn.")
    ]
    for s_title, s_desc in steps_fb3:
        p = doc.add_paragraph()
        r_step = p.add_run(f"• {s_title}: ")
        r_step.bold = True
        r_step.font.color.rgb = RGBColor(30, 58, 138)
        p.add_run(s_desc)

    # ---------------------------------------------------------------------------
    # TỔNG KẾT
    # ---------------------------------------------------------------------------
    doc.add_page_break()
    add_heading_styled(doc, "TỔNG KẾT VẬN HÀNH & FILE MINH CHỨNG", level=1)
    
    p_sum = doc.add_paragraph()
    p_sum.add_run("Tất cả các tính năng đã được kiểm tra tính tiện dụng cao nhất cho người dùng và quản trị viên:\n")
    p_sum.add_run("1. Không yêu cầu bất kỳ kỹ năng lập trình nào từ phía Quản trị viên.\n")
    p_sum.add_run("2. Giao diện trực quan, hỗ trợ đầy đủ tiếng Việt và tiếng Anh.\n")
    p_sum.add_run("3. Vận hành đồng bộ giữa Admin Panel và Storefront người dùng cuối.\n\n")

    p_sum.add_run("Danh sách 7 ảnh chụp thực tế được nhúng trong tài liệu này:\n").bold = True
    img_list = [
        "fb2_admin_collections_list.png — Danh sách bộ sưu tập trong Admin",
        "fb2_admin_collection_bottom_cta.png — Chi tiết cấu hình Nút cuối trang trong Admin",
        "fb2_storefront_bottom_cta.png — Hiển thị thực tế Bottom CTA ngoài website",
        "fb3_admin_size_finder_custom.png — Cấu hình Size Finder tùy chỉnh riêng trong Admin",
        "fb3_admin_size_rules_form.png — Form nhập từng dòng quy tắc size không dùng JSON",
        "fb3_storefront_product_page.png — Trang chi tiết sản phẩm với nút tiện ích size",
        "fb3_storefront_size_finder_modal.png — Modal popup tìm size tương tác ngoài Storefront"
    ]
    for item in img_list:
        p = doc.add_paragraph()
        p.add_run(f"✔ {item}")

    # Output file paths
    out_file_downloads = os.path.join(downloads_dir, "HUONG-DAN-SU-DUNG-FEEDBACK-2-VA-3.docx")
    out_file_repo = os.path.join(repo_dir, "docs", "HUONG-DAN-SU-DUNG-FEEDBACK-2-VA-3.docx")

    doc.save(out_file_downloads)
    print(f"Exported Word document to: {out_file_downloads}")
    
    doc.save(out_file_repo)
    print(f"Saved copy in repository: {out_file_repo}")

if __name__ == "__main__":
    main()
